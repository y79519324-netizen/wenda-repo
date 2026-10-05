"""
文件解析器 - 支持多种文件格式的文本提取

支持的格式:
- .txt: 纯文本文件
- .docx: Word 文档
- .pdf: PDF 文档
"""

import io
from app.core.logger import logger


class FileParser:
    """文件解析器，支持多种格式的文本提取"""
    
    @staticmethod
    def parse_file(file_content: bytes, filename: str) -> str:
        """
        根据文件扩展名选择合适的解析器
        
        Args:
            file_content: 文件二进制内容
            filename: 文件名（包含扩展名）
            
        Returns:
            提取的文本内容
        """
        # 获取文件扩展名（转为小写）
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        
        parser_map = {
            'txt': FileParser._parse_txt,
            'docx': FileParser._parse_docx,
            'pdf': FileParser._parse_pdf,
            'md': FileParser._parse_txt,      # Markdown 也按文本处理
            'csv': FileParser._parse_txt,    # CSV 也按文本处理
            'json': FileParser._parse_txt,   # JSON 也按文本处理
        }
        
        parser = parser_map.get(ext)
        if parser:
            try:
                return parser(file_content)
            except Exception as e:
                logger.error(f"[FileParser] 解析文件失败: {e}")
                raise ValueError(f"解析文件失败: {str(e)}")
        else:
            # 对于未知格式，尝试按文本处理
            logger.warning(f"[FileParser] 未知文件格式 .{ext}，尝试按文本处理")
            try:
                return FileParser._parse_txt(file_content)
            except:
                raise ValueError(f"不支持的文件格式: .{ext}")
    
    @staticmethod
    def _parse_txt(file_content: bytes) -> str:
        """解析纯文本文件"""
        try:
            # 尝试 UTF-8 解码
            return file_content.decode('utf-8')
        except UnicodeDecodeError:
            try:
                # 尝试 GBK 解码（中文 Windows 常用编码）
                return file_content.decode('gbk')
            except UnicodeDecodeError:
                # 使用 errors='ignore' 忽略无法解码的字符
                return file_content.decode('utf-8', errors='ignore')
    
    @staticmethod
    def _parse_docx(file_content: bytes) -> str:
        """解析 Word 文档 (.docx)"""
        try:
            from docx import Document
            
            # 使用 BytesIO 包装二进制内容
            doc = Document(io.BytesIO(file_content))
            
            # 提取所有段落的文本
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():  # 跳过空段落
                    paragraphs.append(para.text)
            
            # 提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_texts = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_texts.append(cell.text.strip())
                    if row_texts:
                        paragraphs.append(' | '.join(row_texts))
            
            logger.info(f"[FileParser] Word 文档解析成功，共 {len(paragraphs)} 个段落")
            return '\n'.join(paragraphs)
            
        except ImportError:
            raise ImportError("请安装 python-docx 库: pip install python-docx")
        except Exception as e:
            raise ValueError(f"解析 Word 文档失败: {str(e)}")
    
    @staticmethod
    def _parse_pdf(file_content: bytes) -> str:
        """解析 PDF 文档"""
        try:
            from pypdf import PdfReader
            
            # 使用 BytesIO 包装二进制内容
            pdf = PdfReader(io.BytesIO(file_content))
            
            # 提取每一页的文本
            pages_text = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text and text.strip():
                    pages_text.append(text.strip())
            
            logger.info(f"[FileParser] PDF 文档解析成功，共 {len(pdf.pages)} 页")
            return '\n\n'.join(pages_text)
            
        except ImportError:
            raise ImportError("请安装 pypdf 库: pip install pypdf")
        except Exception as e:
            raise ValueError(f"解析 PDF 文档失败: {str(e)}")
    
    @staticmethod
    def get_supported_extensions() -> list:
        """获取支持的文件扩展名列表"""
        return ['txt', 'docx', 'pdf', 'md', 'csv', 'json']
    
    @staticmethod
    def get_supported_type_hint() -> str:
        """获取支持的文件类型提示"""
        return ".txt, .docx, .pdf, .md, .csv, .json"
